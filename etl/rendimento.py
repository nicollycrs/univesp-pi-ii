"""
Leitor das planilhas de Taxas de Rendimento Escolar do INEP.

Formato conferido em `tx_rend_escolas_2023.xlsx` no dia 23/09/2026:

  aba              ESCOLAS
  linhas 1 a 8     titulos e cabecalho hierarquico em celulas mescladas
  linha 9          codigos legiveis por maquina (NU_ANO_CENSO, CO_ENTIDADE, ...)
  linha 10 em diante  dados, com "--" onde o indicador nao se aplica
  volume           cerca de 128 mil escolas, 63 colunas

A leitura usa openpyxl em modo somente leitura e coleta apenas as colunas
necessarias, em vez de carregar a planilha inteira com pandas. O arquivo tem
39 MB e mais de 8 milhoes de celulas; ler tudo gastaria memoria sem proveito.

Download: https://download.inep.gov.br/informacoes_estatisticas/indicadores_educacionais/{ano}/tx_rend_escolas_{ano}.zip
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd

from etl.transformar import mapa_etapa, normalizar_nome

# Linha do cabecalho legivel por maquina, contando a partir de 1.
LINHA_CODIGOS = 9
ABA = "ESCOLAS"

# Colunas de identificacao sempre lidas.
IDENTIFICACAO = (
    "CO_ENTIDADE",
    "NO_ENTIDADE",
    "SG_UF",
    "CO_MUNICIPIO",
    "NO_DEPENDENCIA",
    "NO_CATEGORIA",
    "NU_ANO_CENSO",
)


class FormatoInesperado(RuntimeError):
    """A planilha nao tem a estrutura esperada. Provavel mudanca de edicao."""


def _indices_das_colunas(ws, etapa: str) -> tuple[dict[str, int], list[str]]:
    """
    Localiza as colunas pelos codigos da linha 9, em vez de por posicao fixa.

    Posicao fixa quebraria silenciosamente se o INEP inserisse uma coluna. Buscar
    pelo codigo falha alto, com mensagem clara, que e o comportamento desejado.
    """
    codigos = None
    for i, linha in enumerate(ws.iter_rows(min_row=1, max_row=LINHA_CODIGOS, values_only=True), 1):
        if i == LINHA_CODIGOS:
            codigos = [("" if v is None else str(v).strip()) for v in linha]
    if not codigos:
        raise FormatoInesperado(
            f"Nao foi possivel ler a linha {LINHA_CODIGOS} da aba {ABA}."
        )

    posicao = {c: i for i, c in enumerate(codigos) if c}

    alvo = dict(mapa_etapa(etapa))
    desejadas: dict[str, int] = {}
    faltando: list[str] = []

    for codigo in IDENTIFICACAO:
        if codigo in posicao:
            desejadas[codigo] = posicao[codigo]
        else:
            faltando.append(codigo)

    for codigo, indice in posicao.items():
        if normalizar_nome(codigo) in alvo:
            desejadas[codigo] = indice

    esperadas_da_etapa = len(alvo)
    achadas_da_etapa = sum(1 for c in desejadas if normalizar_nome(c) in alvo)
    if achadas_da_etapa < esperadas_da_etapa:
        faltando.append(
            f"colunas de indicador da etapa {etapa!r} "
            f"(esperadas {esperadas_da_etapa}, achadas {achadas_da_etapa})"
        )

    return desejadas, faltando


def ler(caminho: str | Path, etapa: str = "MED", limite: int | None = None) -> pd.DataFrame:
    """
    Le a planilha e devolve um dataframe estreito, ainda cru.

    A limpeza fica com `etl.transformar.transformar()`: aqui so se extrai o
    recorte de colunas. `limite` serve para inspecionar o arquivo rapidamente.
    """
    caminho = Path(caminho)
    if not caminho.exists():
        raise FileNotFoundError(f"Planilha nao encontrada: {caminho}")

    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    try:
        if ABA not in wb.sheetnames:
            raise FormatoInesperado(
                f"A planilha nao tem a aba {ABA!r}. Abas encontradas: {wb.sheetnames}"
            )
        ws = wb[ABA]

        desejadas, faltando = _indices_das_colunas(ws, etapa)
        if faltando:
            raise FormatoInesperado(
                "A planilha nao tem as colunas esperadas: "
                + ", ".join(faltando)
                + ". O INEP provavelmente mudou o layout; ajuste etl/rendimento.py "
                "e registre em docs/FONTES_DE_DADOS.md."
            )

        nomes = list(desejadas)
        indices = [desejadas[n] for n in nomes]

        registros = []
        for i, linha in enumerate(ws.iter_rows(min_row=LINHA_CODIGOS + 1, values_only=True)):
            if limite is not None and i >= limite:
                break
            # Linhas de rodape do INEP vem com a chave vazia.
            if linha[desejadas["CO_ENTIDADE"]] in (None, ""):
                continue
            registros.append([linha[j] for j in indices])
    finally:
        wb.close()

    return pd.DataFrame(registros, columns=nomes)


def caminho_padrao(ano: int, raiz: str | Path = "dados_brutos") -> Path:
    """Caminho onde o zip do INEP extrai a planilha daquele ano."""
    raiz = Path(raiz)
    return raiz / f"tx_rend_escolas_{ano}" / f"tx_rend_escolas_{ano}.xlsx"
